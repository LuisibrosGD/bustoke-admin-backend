import io
from datetime import datetime
from typing import Any, Optional

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BadRequestException, NotFoundException


async def reporte_ventas(
    db: AsyncSession,
    id_agencia: Optional[int] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    id_ruta: Optional[str] = None,
    estado_pago: Optional[str] = None,
    metodo_pago: Optional[str] = None,
    canal_venta: Optional[str] = None,
) -> list[dict[str, Any]]:
    filters = []
    params: dict[str, Any] = {}

    if id_agencia:
        filters.append("a.id_agencia = :id_agencia")
        params["id_agencia"] = id_agencia
    if fecha_inicio:
        filters.append("b.fecha_emision >= :fecha_inicio")
        params["fecha_inicio"] = datetime.fromisoformat(fecha_inicio)
    if fecha_fin:
        filters.append("b.fecha_emision <= :fecha_fin")
        params["fecha_fin"] = datetime.fromisoformat(fecha_fin)
    if id_ruta:
        ids = [x.strip() for x in id_ruta.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":ruta_{i}" for i in range(len(ids)))
            filters.append(f"v.id_ruta IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"ruta_{i}"] = int(val)
    if canal_venta:
        ids = [x.strip() for x in canal_venta.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":canal_{i}" for i in range(len(ids)))
            filters.append(f"b.canal IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"canal_{i}"] = val
    if metodo_pago:
        ids = [x.strip() for x in metodo_pago.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":metodo_{i}" for i in range(len(ids)))
            filters.append(f"pg.metodo IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"metodo_{i}"] = val
    if estado_pago:
        ids = [x.strip() for x in estado_pago.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":epago_{i}" for i in range(len(ids)))
            filters.append(f"pg.estado IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"epago_{i}"] = val

    where_clause = ("WHERE " + " AND ".join(filters)) if filters else ""

    sql = text(f"""
        SELECT
            a.razon_social                                                AS agencia,
            CONCAT(t_o.nombre, ' → ', t_d.nombre)                        AS ruta,
            TO_CHAR(v.fecha_hora_salida, 'YYYY-MM')                       AS fecha_viaje,
            COUNT(DISTINCT b.id_boleto)                                   AS total_boletos,
            SUM(b.precio_final)                                           AS total_ventas,
            COALESCE(
                SUM(b.precio_final) * MAX(cc.porcentaje_comision) / 100,
                0
            )                                                              AS comision
        FROM boletos b
        JOIN viajes v   ON b.id_viaje = v.id_viaje
        JOIN rutas r    ON v.id_ruta  = r.id_ruta
        JOIN terminales t_o ON r.id_terminal_origen = t_o.id_terminal
        JOIN terminales t_d ON r.id_terminal_destino = t_d.id_terminal
        JOIN agencias a ON r.id_agencia = a.id_agencia
        LEFT JOIN pagos pg ON pg.id_boleto = b.id_boleto
        LEFT JOIN configuracion_comisiones cc
            ON cc.id_agencia = a.id_agencia
           AND cc.fecha_inicio <= b.fecha_emision
           AND (cc.fecha_fin IS NULL OR cc.fecha_fin >= b.fecha_emision)
        {where_clause}
        GROUP BY a.razon_social, ruta, fecha_viaje
        ORDER BY fecha_viaje DESC, a.razon_social, ruta
    """)

    result = await db.execute(sql, params)
    rows = result.mappings().all()
    return [
        {
            "agencia": row["agencia"],
            "ruta": row["ruta"],
            "fechaViaje": row["fecha_viaje"],
            "totalBoletos": int(row["total_boletos"]),
            "totalVentas": float(row["total_ventas"]),
            "comision": float(row["comision"]),
        }
        for row in rows
    ]


async def reporte_viajes(
    db: AsyncSession,
    id_agencia: Optional[int] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    id_ruta: Optional[str] = None,
    id_bus: Optional[str] = None,
    estado_viaje: Optional[str] = None,
) -> list[dict[str, Any]]:
    filters = []
    params: dict[str, Any] = {}

    if id_agencia:
        filters.append("r.id_agencia = :id_agencia")
        params["id_agencia"] = id_agencia
    if fecha_inicio:
        filters.append("v.fecha_hora_salida >= :fecha_inicio")
        params["fecha_inicio"] = datetime.fromisoformat(fecha_inicio)
    if fecha_fin:
        filters.append("v.fecha_hora_salida <= :fecha_fin")
        params["fecha_fin"] = datetime.fromisoformat(fecha_fin)
    if id_ruta:
        ids = [x.strip() for x in id_ruta.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":ruta_{i}" for i in range(len(ids)))
            filters.append(f"v.id_ruta IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"ruta_{i}"] = int(val)
    if id_bus:
        ids = [x.strip() for x in id_bus.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":bus_{i}" for i in range(len(ids)))
            filters.append(f"v.id_bus IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"bus_{i}"] = int(val)
    if estado_viaje:
        ids = [x.strip() for x in estado_viaje.split(",") if x.strip()]
        if ids:
            placeholders = ", ".join(f":eviaje_{i}" for i in range(len(ids)))
            filters.append(f"v.estado IN ({placeholders})")
            for i, val in enumerate(ids):
                params[f"eviaje_{i}"] = val

    where_clause = ("WHERE " + " AND ".join(filters)) if filters else ""

    sql = text(f"""
        SELECT
            v.id_viaje,
            CONCAT(t_o.nombre, ' → ', t_d.nombre) AS ruta,
            v.fecha_hora_salida                    AS fecha_salida,
            v.estado,
            COUNT(b.id_boleto)                     AS total_boletos,
            COALESCE(
                ROUND(COUNT(b.id_boleto)::numeric / NULLIF(bus.cantidad_pisos * 50, 0) * 100, 2),
                0
            )                                      AS ocupacion
        FROM viajes v
        JOIN rutas r        ON v.id_ruta              = r.id_ruta
        JOIN terminales t_o ON r.id_terminal_origen   = t_o.id_terminal
        JOIN terminales t_d ON r.id_terminal_destino  = t_d.id_terminal
        JOIN buses bus      ON v.id_bus               = bus.id_bus
        LEFT JOIN boletos b ON b.id_viaje = v.id_viaje AND b.estado = 'activo'
        {where_clause}
        GROUP BY v.id_viaje, ruta, v.fecha_hora_salida, v.estado, bus.cantidad_pisos
        ORDER BY v.fecha_hora_salida DESC
    """)

    result = await db.execute(sql, params)
    rows = result.mappings().all()
    return [
        {
            "idViaje": row["id_viaje"],
            "ruta": row["ruta"],
            "fechaSalida": str(row["fecha_salida"]),
            "estado": row["estado"],
            "totalBoletos": int(row["total_boletos"]),
            "ocupacion": float(row["ocupacion"]),
        }
        for row in rows
    ]


async def reporte_manifiesto_sutran(
    db: AsyncSession, id_viaje: int
) -> list[dict[str, Any]]:
    sql = text("""
        SELECT
            b.id_boleto,
            v.id_viaje,
            CONCAT(p.nombres, ' ', p.apellido_paterno, ' ', p.apellido_materno) AS pasajero,
            p.numero_documento,
            a.numero_asiento                                                      AS asiento,
            b.email_contacto,
            b.precio_final
        FROM boletos b
        JOIN pasajeros p ON b.id_pasajero = p.id_pasajero
        JOIN asientos a  ON b.id_asiento  = a.id_asiento
        JOIN viajes v    ON b.id_viaje    = v.id_viaje
        WHERE v.id_viaje = :id_viaje
          AND b.estado   = 'activo'
        ORDER BY a.numero_asiento
    """)
    result = await db.execute(sql, {"id_viaje": id_viaje})
    rows = result.mappings().all()
    return [
        {
            "idViaje": row["id_viaje"],
            "idBoleto": row["id_boleto"],
            "pasajero": row["pasajero"],
            "numeroDocumento": row["numero_documento"],
            "asiento": row["asiento"],
            "emailContacto": row["email_contacto"],
            "precioFinal": float(row["precio_final"]),
        }
        for row in rows
    ]


async def reporte_financiero(
    db: AsyncSession,
    id_agencia: Optional[int] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
) -> list[dict[str, Any]]:
    filters = []
    params: dict[str, Any] = {}

    if id_agencia:
        filters.append("r.id_agencia = :id_agencia")
        params["id_agencia"] = id_agencia
    if fecha_inicio:
        filters.append("b.fecha_emision >= :fecha_inicio")
        params["fecha_inicio"] = datetime.fromisoformat(fecha_inicio)
    if fecha_fin:
        filters.append("b.fecha_emision <= :fecha_fin")
        params["fecha_fin"] = datetime.fromisoformat(fecha_fin)

    where_clause = ("WHERE " + " AND ".join(filters)) if filters else ""

    sql = text(f"""
        SELECT
            a.razon_social                                    AS agencia,
            TO_CHAR(b.fecha_emision, 'YYYY-MM')               AS periodo,
            COUNT(DISTINCT b.id_boleto)                       AS total_boletos,
            SUM(b.precio_final)                               AS total_ventas,
            COALESCE(
                SUM(b.precio_final) * MAX(cc.porcentaje_comision) / 100,
                0
            )                                                  AS comision,
            COALESCE(
                SUM(b.precio_final)
                - SUM(b.precio_final) * MAX(cc.porcentaje_comision) / 100,
                SUM(b.precio_final)
            )                                                  AS neto_transferir
        FROM boletos b
        JOIN viajes v   ON b.id_viaje = v.id_viaje
        JOIN rutas r    ON v.id_ruta  = r.id_ruta
        JOIN agencias a ON r.id_agencia = a.id_agencia
        LEFT JOIN configuracion_comisiones cc
            ON cc.id_agencia = a.id_agencia
           AND cc.fecha_inicio <= b.fecha_emision
           AND (cc.fecha_fin IS NULL OR cc.fecha_fin >= b.fecha_emision)
        {where_clause}
        GROUP BY a.razon_social, periodo
        ORDER BY periodo DESC, a.razon_social
    """)

    result = await db.execute(sql, params)
    rows = result.mappings().all()
    return [
        {
            "agencia": row["agencia"],
            "periodo": row["periodo"],
            "totalBoletos": int(row["total_boletos"]),
            "totalVentas": float(row["total_ventas"]),
            "comision": float(row["comision"]),
            "netoTransferir": float(row["neto_transferir"]),
        }
        for row in rows
    ]


async def export_excel(data: list[dict[str, Any]], sheet_name: str = "Reporte") -> bytes:
    """Genera un archivo Excel en memoria y devuelve los bytes."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not data:
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    headers = list(data[0].keys())
    ws.append(headers)
    for row in data:
        ws.append([row.get(h) for h in headers])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
